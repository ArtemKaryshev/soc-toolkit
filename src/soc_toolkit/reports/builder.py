"""Генерация отчёта по алерту: Markdown (шаблоны Jinja2), HTML и DOCX.

Поля повторяют шаблон TryHackMe SOC Simulator: Time of activity, List of Affected Entities,
Reason for Classifying as TP/FP, Reason for Escalating, Recommended Remediation Actions,
List of Attack Indicators. Дополнительно: таймлайн, MITRE ATT&CK, результаты репутации.
Свои шаблоны report_<lang>.md.j2 можно положить в папку report.custom_templates_dir.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Any

import markdown
import yaml
from jinja2 import ChoiceLoader, FileSystemLoader, StrictUndefined
from jinja2.sandbox import SandboxedEnvironment

from ..alerts import AlertSummary
from ..ioc import IOCType
from ..reference import technique

HERE = Path(__file__).resolve().parent
LANGS = ("ru", "en")

ENTITY_LABELS = {
    "ru": {"hosts": "Хосты", "users": "Пользователи", "src_ips": "IP-источники", "dest_ips": "IP-назначения",
           "processes": "Процессы", "parent_processes": "Родительские процессы", "files": "Файлы",
           "hashes": "Хэши", "urls": "URL", "domains": "Домены"},
    "en": {"hosts": "Hosts", "users": "Users", "src_ips": "Source IPs", "dest_ips": "Destination IPs",
           "processes": "Processes", "parent_processes": "Parent processes", "files": "Files",
           "hashes": "Hashes", "urls": "URLs", "domains": "Domains"},
}


@lru_cache
def remediation_catalog() -> dict[str, Any]:
    return yaml.safe_load((HERE / "remediation.yaml").read_text(encoding="utf-8"))["categories"]


def categories() -> list[str]:
    return [c for c in remediation_catalog() if c != "fp"]


def category_title(category: str, lang: str = "ru") -> str:
    return remediation_catalog().get(category, {}).get("title", {}).get(lang, category)


def default_remediation(category: str, verdict: str, lang: str = "ru") -> list[str]:
    key = "fp" if verdict.upper() == "FP" else category
    return list(remediation_catalog().get(key, remediation_catalog()["generic"])[lang])


@lru_cache
def _system_binaries() -> frozenset[str]:
    """Имена системных утилит Windows: сами по себе не индикаторы атаки."""
    from ..reference import lolbas

    names = {e["name"].lower() for e in lolbas()}
    names.update({"cmd.exe", "powershell.exe", "pwsh.exe", "explorer.exe", "svchost.exe", "services.exe",
                  "lsass.exe", "winword.exe", "excel.exe", "outlook.exe", "powerpnt.exe", "chrome.exe",
                  "msedge.exe", "firefox.exe", "conhost.exe", "taskhostw.exe", "wmiprvse.exe", "notepad.exe"})
    return frozenset(names)


@dataclass
class ReportInput:
    verdict: str                          # TP / FP
    reason: str = ""
    escalate: bool = False
    escalation_reason: str = ""
    category: str = "generic"
    remediation: list[str] | None = None  # None: взять типовые для категории
    notes: str = ""
    lang: str = "ru"
    analyst: str = ""
    case_id: int | None = None
    defang: bool = True
    reputation: dict[str, str] = field(default_factory=dict)   # значение IOC -> сводный вердикт


def _context(alert: AlertSummary, inp: ReportInput) -> dict[str, Any]:
    lang = inp.lang if inp.lang in LANGS else "ru"
    is_tp = inp.verdict.upper() == "TP"
    entities = [(ENTITY_LABELS[lang][k], ", ".join(f"`{v}`" for v in getattr(alert, k)))
                for k in ENTITY_LABELS[lang] if getattr(alert, k)]

    indicators = []
    for i in alert.iocs:
        # Индикаторы атаки: внешние сетевые, хэши, вредоносные файлы. Внутренние адреса есть в entities.
        if i.type in (IOCType.MITRE, IOCType.CVE, IOCType.WIN_PATH, IOCType.REG_KEY):
            continue
        if i.scope not in ("", "public") or "allowlisted" in i.tags:
            continue
        if i.type == IOCType.FILENAME and i.value.lower() in _system_binaries():
            continue
        indicators.append({"type": i.type.value, "value": i.defanged if inp.defang else i.value,
                           "verdict": inp.reputation.get(i.value, "")})

    mitre = []
    for h in alert.mitre:
        t = technique(h.technique)
        reason = h.reason if lang == "ru" else ("listed in the alert" if h.explicit else f"heuristic: {h.reason}")
        url = t["url"] if t else f"https://attack.mitre.org/techniques/{h.technique.replace('.', '/')}/"
        mitre.append({"id": h.technique, "name": t["full_name"] if t else "", "reason": reason, "url": url})

    remediation = inp.remediation if inp.remediation is not None else default_remediation(inp.category, inp.verdict,
                                                                                         lang)
    return {
        "a": alert,
        "ctx": {
            "is_tp": is_tp, "reason": inp.reason.strip(), "escalate": inp.escalate,
            "escalation_reason": inp.escalation_reason.strip(), "remediation": [r for r in remediation if r.strip()],
            "category_title": category_title("fp" if not is_tp else inp.category, lang),
            "entities": entities, "indicators": indicators, "mitre": mitre, "notes": inp.notes.strip(),
            "analyst": inp.analyst, "case_id": inp.case_id, "created": datetime.now().strftime("%Y-%m-%d %H:%M"),
        },
    }


def _env(custom_dir: str = "") -> SandboxedEnvironment:
    loaders = []
    if custom_dir and Path(custom_dir).expanduser().is_dir():
        loaders.append(FileSystemLoader(str(Path(custom_dir).expanduser())))
    loaders.append(FileSystemLoader(str(HERE / "templates")))
    return SandboxedEnvironment(loader=ChoiceLoader(loaders), undefined=StrictUndefined, trim_blocks=True,
                                lstrip_blocks=True, autoescape=False)


def render_markdown(alert: AlertSummary, inp: ReportInput, custom_dir: str = "") -> str:
    lang = inp.lang if inp.lang in LANGS else "ru"
    text = _env(custom_dir).get_template(f"report_{lang}.md.j2").render(**_context(alert, inp))
    # убираем лишние пустые строки, которые оставляют условные блоки
    lines: list[str] = []
    for line in text.splitlines():
        if line.strip() == "" and lines and lines[-1].strip() == "":
            continue
        lines.append(line.rstrip())
    return "\n".join(lines).strip() + "\n"


HTML_SHELL = """<!doctype html>
<html lang="{lang}"><head><meta charset="utf-8"><title>{title}</title>
<style>
body {{ font-family: -apple-system, Segoe UI, Roboto, Arial, sans-serif; max-width: 960px; margin: 2rem auto;
       padding: 0 1rem; color: #1d2433; line-height: 1.5; }}
h1 {{ border-bottom: 3px solid #c0392b; padding-bottom: .3rem; }}
h2 {{ margin-top: 2rem; color: #2c3e50; border-bottom: 1px solid #dde; padding-bottom: .2rem; }}
table {{ border-collapse: collapse; width: 100%; margin: .5rem 0; }}
td, th {{ border: 1px solid #ccd; padding: .35rem .6rem; text-align: left; vertical-align: top; }}
th {{ background: #f3f5f9; }}
code, pre {{ background: #f4f4f7; border-radius: 4px; font-family: Consolas, monospace; font-size: .9em; }}
code {{ padding: .1rem .3rem; word-break: break-all; }}
pre {{ padding: .6rem; white-space: pre-wrap; word-break: break-all; }}
</style></head><body>
{body}
</body></html>"""


def render_html(md_text: str, lang: str = "ru") -> str:
    body = markdown.markdown(md_text, extensions=["tables", "fenced_code"])
    title = md_text.splitlines()[0].lstrip("# ").strip() if md_text else "Report"
    return HTML_SHELL.format(lang=lang, title=title, body=body)


def render_docx(md_text: str) -> bytes:
    """DOCX из Markdown отчёта: заголовки, списки, таблицы, код. Без внешних конвертеров."""
    import re

    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    def add_runs(par: Any, text: str) -> None:
        # **жирный**, `код`, [текст](ссылка)
        text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if part.startswith("**") and part.endswith("**"):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("`") and part.endswith("`"):
                run = par.add_run(part[1:-1])
                run.font.name = "Consolas"
            elif part:
                par.add_run(part)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.fullmatch(r"\|[\s|:-]+\|?", lines[i]):
                    rows.append([c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", lines[i])[1:-1]])
                i += 1
            i -= 1
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=ncols)
                table.style = "Table Grid"
                for r in rows:
                    cells = table.add_row().cells
                    for c, text in zip(cells, r, strict=False):
                        add_runs(c.paragraphs[0], text)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif line.strip():
            add_runs(doc.add_paragraph(), line)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
